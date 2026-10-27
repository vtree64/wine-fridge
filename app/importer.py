"""Shared spreadsheet import logic (used by the CLI importer and the API)."""

import datetime

from openpyxl import load_workbook

from . import db

# Excel's epoch for serial date numbers
EXCEL_EPOCH = datetime.datetime(1899, 12, 30)


def cell_text(value, name: str) -> str:
    """Normalize a spreadsheet cell to the string we store."""
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, float):
        if name == "Date Added":
            return (EXCEL_EPOCH + datetime.timedelta(days=value)).strftime("%Y-%m-%d %H:%M")
        if name == "Vintage" and value.is_integer():
            return str(int(value))
        return str(value)
    return str(value).strip()


def import_from_xlsx(xlsx_path: str) -> tuple[int, int, int]:
    """Import bottles from a Wine Inventory spreadsheet.

    Idempotent: bottles whose Bottle ID already exists are skipped.
    Returns (added, skipped, total).
    """
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Spreadsheet is empty")

    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    idx = {name: i for i, name in enumerate(header)}
    missing = [c for c in ("Bottle ID", "Wine", "Shelf") if c not in idx]
    if missing:
        raise ValueError(f"Missing required column(s): {', '.join(missing)}. Found: {header}")

    def cell(row, name):
        i = idx.get(name)
        if i is None or i >= len(row):
            return ""
        return cell_text(row[i], name)

    db.init_db()

    added = 0
    skipped = 0
    for row in rows[1:]:
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue
        bottle_id = cell(row, "Bottle ID")
        wine = cell(row, "Wine")
        if not bottle_id or not wine:
            skipped += 1
            continue
        if db.get_bottle_row(bottle_id) is not None:
            skipped += 1
            continue
        db.add_bottle(
            {
                "wine": wine,
                "variety": cell(row, "Variety"),
                "vintage": cell(row, "Vintage"),
                "vineyard_producer": cell(row, "Vineyard / Producer"),
                "shelf": cell(row, "Shelf"),
                "notes": cell(row, "Notes"),
            },
            bottle_id=bottle_id,
            date_added=cell(row, "Date Added")
            or datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
        )
        added += 1

    return added, skipped, db.count_bottles()
